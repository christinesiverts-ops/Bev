"""Build Chain_Audit_Tool.xlsx from tools/program_data.py.

    python tools/build_workbook.py [output.xlsx]

Re-run after editing program_data.py. Once the team is live in the workbook,
edit the workbook directly instead (re-running overwrites the Visit Log).
"""
import os
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import program_data as pdat  # noqa: E402

AS_OF = "Plan data as of Non-Alc Chain Document 9/8/2026 (+ HW Soda Hot Sheet 8/25/2025 for accounts missing from it)"

FONT = "Arial"
NAVY = "1F2A44"
TEAL = "0F6B78"
BROWN = "8A5A19"
WINE = "7A1F2B"
SLATE = "4A5563"
LIGHT = "F3F4F6"
INPUT_FILL = "FFF7D6"   # rep entry cells
BLUE = "0000FF"         # manager-maintained inputs
BRAND_COLORS = {"ZOA": "E4F3F5", "Henry's": "F7E9E4", "Fever-Tree": "F4EEDF", "Naked Life": "E6F2E6"}

thin = Side(style="thin", color="D1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

CP_FIRST = 6                              # first data row on Chain Program / Promo Calendar / Visit Log
CP_LAST = int(os.environ.get("CP_LAST", 120))    # room to add programs
PC_LAST = int(os.environ.get("PC_LAST", 400))
VL_LAST = int(os.environ.get("VL_LAST", 1005))   # 1,000 visit rows (env overrides are for quick tests)
HIST_ROWS = 20


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def title(ws, text, sub, width_cols):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 16)
    ws["A2"] = sub
    ws["A2"].font = f(False, SLATE, 9, italic=True)
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24


def group_band(ws, row, groups):
    """groups: list of (label, first_col, last_col, color)"""
    for label, c1, c2, color in groups:
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = ws.cell(row=row, column=c1, value=label)
        cell.font = f(True, "FFFFFF", 10)
        cell.fill = fill(color)
        cell.alignment = CENTER
        for c in range(c1, c2 + 1):
            ws.cell(row=row, column=c).fill = fill(color)
    ws.row_dimensions[row].height = 18


def headers(ws, row, cols, color_for_col):
    for i, (name, width) in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = f(True, "FFFFFF", 9)
        cell.fill = fill(color_for_col(i))
        cell.alignment = CENTER
        cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 42


def add_list_validation(ws, formula, rng, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Pick a value from the list."
    dv.errorTitle = "Not on the list"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ----------------------------------------------------------------------------------------------
wb = Workbook()
ws_readme = wb.active
ws_readme.title = "README"
ws_dash = wb.create_sheet("Dashboard")
ws_cp = wb.create_sheet("Chain Program")
ws_pc = wb.create_sheet("Promo Calendar")
ws_vl = wb.create_sheet("Visit Log")
ws_sh = wb.create_sheet("Store History")
ws_dr = wb.create_sheet("Data Review")
ws_ls = wb.create_sheet("Lists")

# ---------------------------------------------------------------- Lists
ws_ls.sheet_view.showGridLines = False
ws_ls["A1"] = "Dropdown lists (manager maintains)"
ws_ls["A1"].font = f(True, NAVY, 14)
list_defs = [
    ("Team Roster", ["EXAMPLE - delete me"] + pdat.ROSTER + [""] * 23),
    ("Brand", pdat.BRANDS),
    ("SKU / Pack", pdat.SKU_LIST),
    ("Discrepancy Type", pdat.DISCREPANCY_TYPES),
    ("Status", pdat.STATUSES),
    ("Y/N", pdat.YN),
    ("Display", pdat.DISPLAY_YN),
]
LIST_RANGES = {}
for ci, (name, vals) in enumerate(list_defs, start=1):
    col = get_column_letter(ci)
    h = ws_ls.cell(row=3, column=ci, value=name)
    h.font = f(True, "FFFFFF", 10)
    h.fill = fill(NAVY)
    h.alignment = CENTER
    ws_ls.column_dimensions[col].width = 34 if ci in (3, 4) else 22
    for ri, v in enumerate(vals, start=4):
        c = ws_ls.cell(row=ri, column=ci, value=v if v != "" else None)
        c.font = f(color=BLUE)
        c.border = BOX
    last = 3 + max(len(vals), 1)
    LIST_RANGES[name] = f"Lists!${col}$4:${col}${last}"
ws_ls["A31"] = "Add team members in the Team Roster column (blank rows are already in the dropdown range)."
ws_ls["A31"].font = f(italic=True, color=SLATE, size=9)

# ---------------------------------------------------------------- Chain Program
CP_COLS = [
    ("Program ID", 34), ("Brand", 11), ("Chain / Banner", 20), ("Account / Division", 24), ("Market", 18),
    ("Outlets by Market Unit", 26), ("Total Outlets", 9),
    ("Authorized SKUs / Flavors", 30), ("Case Recommendation", 20),
    ("Shelf Location / Display Callout", 30),
    ("Case Cost / Frontline", 20), ("EDV / SRP", 18), ("Base Shelf Price ($)", 10), ("Promo Timing", 20),
    ("Current Promo (today)", 24), ("Expected Shelf Price Today ($)", 12), ("Next Promo Start", 11),
    ("Execution Priority / Callouts", 46), ("Owner", 16), ("Source", 18), ("Data Status", 14),
    ("Last Verified", 11), ("Visits (30 days)", 9), ("Open Issues", 9), ("Last Visit", 11),
]
CP_GROUPS = [("CHANNEL", 1, 7, NAVY), ("AVAILABILITY", 8, 9, TEAL), ("MERCHANDISING", 10, 10, BROWN),
             ("PRICING / PROMOTION", 11, 17, WINE), ("EXECUTION & ACCOUNTABILITY", 18, 25, SLATE)]


def group_color(groups):
    def inner(i):
        for _, a, b, c in groups:
            if a <= i <= b:
                return c
        return SLATE
    return inner


title(ws_cp, "CHAIN PROGRAM - HOT SHEET", AS_OF + ".  Blue text = manager inputs.  Grey columns = formulas (do not type).", 25)
ws_cp["A3"] = "Today:"
ws_cp["A3"].font = f(True)
ws_cp["B3"] = "=TODAY()"
ws_cp["B3"].number_format = "m/d/yyyy"
ws_cp["B3"].font = f(True)
group_band(ws_cp, 4, CP_GROUPS)
headers(ws_cp, 5, CP_COLS, group_color(CP_GROUPS))

PCR = f"'Promo Calendar'!$A${CP_FIRST}:$A${PC_LAST}"   # program id
PCS = f"'Promo Calendar'!$D${CP_FIRST}:$D${PC_LAST}"   # sku
PCF = f"'Promo Calendar'!$F${CP_FIRST}:$F${PC_LAST}"   # offer
PCG = f"'Promo Calendar'!$G${CP_FIRST}:$G${PC_LAST}"   # unit retail
PCI = f"'Promo Calendar'!$I${CP_FIRST}:$I${PC_LAST}"   # start
PCJ = f"'Promo Calendar'!$J${CP_FIRST}:$J${PC_LAST}"   # end
PCL = f"'Promo Calendar'!$L${CP_FIRST}:$L${PC_LAST}"   # status
PCO = f"'Promo Calendar'!$O${CP_FIRST}:$O${PC_LAST}"   # active key
VLA = f"'Visit Log'!$A${CP_FIRST}:$A${VL_LAST}"
VLB = f"'Visit Log'!$B${CP_FIRST}:$B${VL_LAST}"
VLC = f"'Visit Log'!$C${CP_FIRST}:$C${VL_LAST}"
VLD = f"'Visit Log'!$D${CP_FIRST}:$D${VL_LAST}"
VLJ = f"'Visit Log'!$J${CP_FIRST}:$J${VL_LAST}"
VLO = f"'Visit Log'!$O${CP_FIRST}:$O${VL_LAST}"
VLR = f"'Visit Log'!$R${CP_FIRST}:$R${VL_LAST}"
VLS = f"'Visit Log'!$S${CP_FIRST}:$S${VL_LAST}"
VLT = f"'Visit Log'!$T${CP_FIRST}:$T${VL_LAST}"
CPA = f"'Chain Program'!$A${CP_FIRST}:$A${CP_LAST}"


def cp_formulas(r):
    a = f"$A{r}"
    active_n = f'COUNTIFS({PCR},{a},{PCL},"Active")'
    retail_n = f'COUNTIFS({PCR},{a},{PCL},"Active",{PCG},">0")'
    return {
        15: f'=IF({a}="","",IF({active_n}=0,"None active",INDEX({PCF},MATCH({a},{PCO},0))'
            f'&IF({active_n}>1," (+"&({active_n}-1)&" more)","")))',
        16: f'=IF({a}="","",IF({retail_n}>0,_xlfn.MINIFS({PCG},{PCR},{a},{PCL},"Active",{PCG},">0"),'
            f'IF(M{r}="","Not set",M{r})))',
        17: f'=IF({a}="","",IF(COUNTIFS({PCR},{a},{PCI},">"&TODAY())=0,"None scheduled",'
            f'_xlfn.MINIFS({PCI},{PCR},{a},{PCI},">"&TODAY())))',
        23: f'=IF({a}="","",COUNTIFS({VLC},{a},{VLA},">="&TODAY()-30))',
        24: f'=IF({a}="","",COUNTIFS({VLC},{a},{VLT},"Open")+COUNTIFS({VLC},{a},{VLT},"In Progress"))',
        25: f'=IF({a}="","",IF(COUNTIFS({VLC},{a})=0,"No visits",_xlfn.MAXIFS({VLA},{VLC},{a})))',
    }


FORMULA_FILL = "EEF0F3"
for i, p in enumerate(pdat.PROGRAMS):
    r = CP_FIRST + i
    vals = [pdat.pid(p["brand"], p["account"]), p["brand"], p["chain"], p["account"], p.get("market", ""),
            p.get("outlets", ""), p.get("total"), p.get("skus", ""), p.get("case_rec", ""), p.get("shelf", ""),
            p.get("case_cost", ""), p.get("srp_text", ""), p.get("base"), p.get("timing", "")]
    for c, v in enumerate(vals, start=1):
        cell = ws_cp.cell(row=r, column=c, value=v if v != "" else None)
        cell.font = f(color=BLUE, bold=(c == 1))
        cell.alignment = WRAP
        cell.border = BOX
    ws_cp.cell(row=r, column=13).number_format = '"$"#,##0.00'
    for c, v in ((18, p.get("priority", "")), (19, p.get("owner", "")), (20, p["source"]), (21, p["status"])):
        cell = ws_cp.cell(row=r, column=c, value=v or None)
        cell.font = f(color=BLUE)
        cell.alignment = WRAP
        cell.border = BOX
    ws_cp.cell(row=r, column=22).border = BOX
    ws_cp.cell(row=r, column=22).number_format = "m/d/yyyy"
    ws_cp.cell(row=r, column=2).fill = fill(BRAND_COLORS[p["brand"]])

for r in range(CP_FIRST, CP_LAST + 1):
    for c, formula in cp_formulas(r).items():
        cell = ws_cp.cell(row=r, column=c, value=formula)
        cell.font = f()
        cell.fill = fill(FORMULA_FILL)
        cell.border = BOX
        cell.alignment = WRAP
    ws_cp.cell(row=r, column=16).number_format = '"$"#,##0.00'
    ws_cp.cell(row=r, column=17).number_format = "m/d/yyyy"
    ws_cp.cell(row=r, column=25).number_format = "m/d/yyyy"

n_prog = len(pdat.PROGRAMS)
ws_cp.freeze_panes = "B6"
ws_cp.auto_filter.ref = f"A5:Y{CP_FIRST + n_prog - 1}"
add_list_validation(ws_cp, f"={LIST_RANGES['Brand']}", f"B{CP_FIRST}:B{CP_LAST}")
dv = DataValidation(type="list", formula1='"Current (2026),Confirm for 2026"', allow_blank=True)
ws_cp.add_data_validation(dv)
dv.add(f"U{CP_FIRST}:U{CP_LAST}")
ws_cp.conditional_formatting.add(
    f"U{CP_FIRST}:U{CP_LAST}", CellIsRule(operator="equal", formula=['"Confirm for 2026"'], fill=fill("FDE68A")))
ws_cp.conditional_formatting.add(
    f"X{CP_FIRST}:X{CP_LAST}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("FECACA")))
ws_cp.conditional_formatting.add(
    f"O{CP_FIRST}:O{CP_LAST}",
    FormulaRule(formula=[f'AND(O{CP_FIRST}<>"",O{CP_FIRST}<>"None active")'], fill=fill("D1FAE5")))
ws_cp.cell(row=5, column=13).comment = Comment(
    "Base Shelf Price = everyday retail the shelf should show when no retail promo is active. "
    "Source per row in the Source column.", "Chain Audit Tool")

# ---------------------------------------------------------------- Promo Calendar
PC_COLS = [
    ("Program ID", 34), ("Brand", 11), ("Account / Division", 24), ("SKU / Pack", 20), ("Offer Type", 13),
    ("Offer", 30), ("Unit Retail ($)", 10), ("Case Cost ($)", 10), ("Start", 11), ("End", 11),
    ("Display Required", 9), ("Status", 11), ("Days Left", 8), ("Source", 20), ("Active Key (helper)", 18),
]
PC_GROUPS = [("PROGRAM", 1, 4, NAVY), ("OFFER", 5, 8, WINE), ("WINDOW", 9, 13, TEAL), ("", 14, 15, SLATE)]
title(ws_pc, "PROMO CALENDAR",
      "One row per promo window. Status updates automatically from today's date. Unit Retail drives the expected "
      "shelf price; leave it blank for case-cost-only rows. SKU 'All' applies to every SKU in the program.", 15)
group_band(ws_pc, 4, PC_GROUPS)
headers(ws_pc, 5, PC_COLS, group_color(PC_GROUPS))
promos = sorted(pdat.PROMOS, key=lambda x: (x[9] != pdat.NA, x[0], x[6]))
for i, pr in enumerate(promos):
    r = CP_FIRST + i
    prog, sku, typ, offer, unit, case, start, end, disp, src = pr
    for c, v in ((1, prog), (4, sku), (5, typ), (6, offer), (7, unit), (8, case), (9, start), (10, end),
                 (11, disp), (14, src)):
        cell = ws_pc.cell(row=r, column=c, value=v)
        cell.font = f(color=BLUE)
        cell.border = BOX
        cell.alignment = WRAP
for r in range(CP_FIRST, PC_LAST + 1):
    forms = {
        2: f'=IF($A{r}="","",IFERROR(INDEX(\'Chain Program\'!$B${CP_FIRST}:$B${CP_LAST},MATCH($A{r},{CPA},0)),"Unknown ID"))',
        3: f'=IF($A{r}="","",IFERROR(INDEX(\'Chain Program\'!$D${CP_FIRST}:$D${CP_LAST},MATCH($A{r},{CPA},0)),"Unknown ID"))',
        12: f'=IF(OR($A{r}="",$I{r}="",$J{r}=""),"",IF(TODAY()<$I{r},"Upcoming",IF(TODAY()>$J{r},"Expired","Active")))',
        13: f'=IF($L{r}="Active",$J{r}-TODAY(),"")',
        15: f'=IF($L{r}="Active",$A{r},"")',
    }
    for c, formula in forms.items():
        cell = ws_pc.cell(row=r, column=c, value=formula)
        cell.font = f(color="6B7280" if c == 15 else "000000")
        cell.fill = fill(FORMULA_FILL)
        cell.border = BOX
    for c in (7, 8):
        ws_pc.cell(row=r, column=c).number_format = '"$"#,##0.00'
        ws_pc.cell(row=r, column=c).border = BOX
    for c in (9, 10):
        ws_pc.cell(row=r, column=c).number_format = "m/d/yyyy"
        ws_pc.cell(row=r, column=c).border = BOX
for c in (1, 4, 5, 6, 11, 14):
    for r in range(CP_FIRST + len(promos), PC_LAST + 1):
        ws_pc.cell(row=r, column=c).border = BOX
ws_pc.freeze_panes = "B6"
ws_pc.auto_filter.ref = f"A5:O{CP_FIRST + len(promos) - 1}"
add_list_validation(ws_pc, f"={CPA}", f"A{CP_FIRST}:A{PC_LAST}")
add_list_validation(ws_pc, f"={LIST_RANGES['SKU / Pack']}", f"D{CP_FIRST}:D{PC_LAST}")
add_list_validation(ws_pc, f"={LIST_RANGES['Y/N']}", f"K{CP_FIRST}:K{PC_LAST}")
for status, color in (("Active", "D1FAE5"), ("Upcoming", "DBEAFE"), ("Expired", "F3F4F6")):
    ws_pc.conditional_formatting.add(
        f"L{CP_FIRST}:L{PC_LAST}", CellIsRule(operator="equal", formula=[f'"{status}"'], fill=fill(color)))
ws_pc.column_dimensions["O"].hidden = False

# ---------------------------------------------------------------- Visit Log
VL_COLS = [
    ("Visit Date", 11), ("Rep", 16), ("Program ID", 34), ("Brand", 11), ("Account / Division", 22),
    ("Store # / Location", 18), ("SKU / Pack Checked", 18), ("Expected Price ($)", 10), ("Observed Shelf Price ($)", 10),
    ("Price Check", 11), ("Shelf Tag Correct?", 8), ("Shelf Location per Planogram?", 10),
    ("All Authorized SKUs In Stock?", 10), ("Display per Callout?", 9), ("Discrepancy Type", 22),
    ("Visit Notes", 40), ("Follow-up Action", 28), ("Follow-up Owner", 16), ("Due Date", 11), ("Status", 11),
    ("Resolved Date", 11), ("Days Open", 8), ("History Match (helper)", 9), ("History Rank (helper)", 9),
]
VL_GROUPS = [("VISIT", 1, 7, NAVY), ("PRICING", 8, 10, WINE), ("SHELF & DISPLAY", 11, 14, BROWN),
             ("DISCREPANCY & NOTES", 15, 17, TEAL), ("ACCOUNTABILITY", 18, 22, SLATE), ("", 23, 24, "9CA3AF")]
title(ws_vl, "VISIT LOG - STORE AUDITS & NOTES",
      "Reps fill the yellow cells only, one row per store / SKU check, newest at the bottom. Grey cells calculate. "
      "Row 6 is an EXAMPLE - delete it before go-live.", 24)
group_band(ws_vl, 4, VL_GROUPS)
headers(ws_vl, 5, VL_COLS, group_color(VL_GROUPS))

INPUT_COLS = {1, 2, 3, 6, 7, 9, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21}
SEL_PROG = "'Store History'!$C$4"
SEL_STORE = "'Store History'!$C$5"
for r in range(CP_FIRST, VL_LAST + 1):
    sku = f'IF($G{r}="","All",$G{r})'
    dt = f'$A{r}'
    def active(sku_expr):
        return (f'{PCR},$C{r},{PCS},{sku_expr},{PCI},"<="&{dt},{PCJ},">="&{dt},{PCG},">0"')
    forms = {
        4: f'=IF($C{r}="","",IFERROR(INDEX(\'Chain Program\'!$B${CP_FIRST}:$B${CP_LAST},MATCH($C{r},{CPA},0)),"Unknown ID"))',
        5: f'=IF($C{r}="","",IFERROR(INDEX(\'Chain Program\'!$D${CP_FIRST}:$D${CP_LAST},MATCH($C{r},{CPA},0)),"Unknown ID"))',
        8: (f'=IF(OR($C{r}="",$A{r}=""),"",IF(COUNTIFS({active(sku)})>0,_xlfn.MINIFS({PCG},{active(sku)}),'
            f'IF(COUNTIFS({active(chr(34) + "All" + chr(34))})>0,_xlfn.MINIFS({PCG},{active(chr(34) + "All" + chr(34))}),'
            f'IFERROR(IF(INDEX(\'Chain Program\'!$M${CP_FIRST}:$M${CP_LAST},MATCH($C{r},{CPA},0))="","Not set",'
            f'INDEX(\'Chain Program\'!$M${CP_FIRST}:$M${CP_LAST},MATCH($C{r},{CPA},0))),"Not set"))))'),
        10: (f'=IF($I{r}="","",IF(NOT(ISNUMBER($H{r})),"No plan price",IF(ABS($I{r}-$H{r})<0.005,"Match",'
             f'IF($I{r}>$H{r},"Over plan","Under plan"))))'),
        22: f'=IF(OR($A{r}="",$T{r}="",$T{r}="No Issue"),"",IF($T{r}="Resolved",IF($U{r}="","",$U{r}-$A{r}),TODAY()-$A{r}))',
        23: f'=IF(AND($C{r}<>"",$C{r}={SEL_PROG},OR({SEL_STORE}="",$F{r}={SEL_STORE})),1,0)',
        24: f'=IF($W{r}=1,SUM($W{r}:$W${VL_LAST}),"")',
    }
    for c in range(1, 25):
        cell = ws_vl.cell(row=r, column=c)
        cell.border = BOX
        if c in forms:
            cell.value = forms[c]
            cell.fill = fill(FORMULA_FILL)
            cell.font = f(color="6B7280" if c >= 23 else "000000")
        elif c in INPUT_COLS:
            cell.fill = fill(INPUT_FILL)
            cell.font = f()
        if c in (16, 17):
            cell.alignment = WRAP
    for c in (1, 19, 21):
        ws_vl.cell(row=r, column=c).number_format = "m/d/yyyy"
    for c in (8, 9):
        ws_vl.cell(row=r, column=c).number_format = '"$"#,##0.00'

# Example row (row 6)
example = {1: date(2026, 9, 24), 2: "EXAMPLE - delete me", 3: pdat.pid("ZOA", "Smart & Final"),
           6: "Example store #000 - Anytown", 7: "Singles (Frosted Grape/Tropical Punch)", 9: 1.79, 11: "N",
           12: "Y", 13: "Y", 14: "No", 15: "Price - over plan",
           16: "Singles ringing $1.79 each; 2 for $3 tag not up. No display built yet.",
           17: "Ask store manager to run the 2/$3 tag; resubmit display authorization.",
           18: "EXAMPLE - delete me", 19: date(2026, 10, 1), 20: "Open"}
for c, v in example.items():
    ws_vl.cell(row=CP_FIRST, column=c, value=v).font = f(italic=True, color="92400E")

ws_vl.freeze_panes = "D6"
ws_vl.auto_filter.ref = f"A5:X{VL_LAST}"
rng = lambda col: f"{col}{CP_FIRST}:{col}{VL_LAST}"  # noqa: E731
add_list_validation(ws_vl, f"={LIST_RANGES['Team Roster']}", rng("B"))
add_list_validation(ws_vl, f"={CPA}", rng("C"), "Pick the brand | account program for this check.")
add_list_validation(ws_vl, f"={LIST_RANGES['SKU / Pack']}", rng("G"))
for col in "KLM":
    add_list_validation(ws_vl, f"={LIST_RANGES['Y/N']}", rng(col))
add_list_validation(ws_vl, f"={LIST_RANGES['Display']}", rng("N"))
add_list_validation(ws_vl, f"={LIST_RANGES['Discrepancy Type']}", rng("O"))
add_list_validation(ws_vl, f"={LIST_RANGES['Team Roster']}", rng("R"))
add_list_validation(ws_vl, f"={LIST_RANGES['Status']}", rng("T"))
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2025,1,1)", allow_blank=True)
dvd.error = "Enter a date (m/d/yyyy)."
ws_vl.add_data_validation(dvd)
dvd.add(rng("A"))
dvn = DataValidation(type="decimal", operator="between", formula1="0", formula2="500", allow_blank=True)
dvn.error = "Enter the shelf price as a number, e.g. 5.99"
ws_vl.add_data_validation(dvn)
dvn.add(rng("I"))
red, green, amber = fill("FECACA"), fill("D1FAE5"), fill("FDE68A")
ws_vl.conditional_formatting.add(rng("J"), CellIsRule(operator="equal", formula=['"Match"'], fill=green))
ws_vl.conditional_formatting.add(rng("J"), CellIsRule(operator="equal", formula=['"Over plan"'], fill=red))
ws_vl.conditional_formatting.add(rng("J"), CellIsRule(operator="equal", formula=['"Under plan"'], fill=amber))
for col in "KLM":
    ws_vl.conditional_formatting.add(rng(col), CellIsRule(operator="equal", formula=['"N"'], fill=red))
ws_vl.conditional_formatting.add(rng("N"), CellIsRule(operator="equal", formula=['"No"'], fill=red))
ws_vl.conditional_formatting.add(rng("T"), CellIsRule(operator="equal", formula=['"Open"'], fill=red))
ws_vl.conditional_formatting.add(rng("T"), CellIsRule(operator="equal", formula=['"In Progress"'], fill=amber))
ws_vl.conditional_formatting.add(rng("T"), CellIsRule(operator="equal", formula=['"Resolved"'], fill=green))
ws_vl.conditional_formatting.add(
    rng("S"), FormulaRule(formula=[f'AND($S{CP_FIRST}<>"",$S{CP_FIRST}<TODAY(),OR($T{CP_FIRST}="Open",$T{CP_FIRST}="In Progress"))'],
                          fill=red, font=Font(name=FONT, bold=True, color="991B1B")))

# ---------------------------------------------------------------- Store History (recall before a visit)
ws_sh.sheet_view.showGridLines = False
ws_sh["A1"] = "STORE HISTORY - CHECK BEFORE YOU WALK IN"
ws_sh["A1"].font = f(True, NAVY, 16)
ws_sh["A2"] = ("Pick a program (and optionally a store #). Shows what the shelf should look like today and the "
               f"last {HIST_ROWS} visit notes, newest first.")
ws_sh["A2"].font = f(italic=True, color=SLATE, size=9)
for r, label in ((4, "Program"), (5, "Store # / Location (optional)")):
    ws_sh.cell(row=r, column=1, value=label).font = f(True)
    ws_sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws_sh.cell(row=r, column=3)
    c.fill = fill(INPUT_FILL)
    c.border = BOX
    c.font = f(True)
    ws_sh.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
ws_sh["C4"] = pdat.pid("ZOA", "Smart & Final")
ws_sh["C5"] = None
add_list_validation(ws_sh, f"={CPA}", "C4")

snap = [("Account / Division", "D"), ("Market", "E"), ("Shelf Location / Display Callout", "J"),
        ("Authorized SKUs / Flavors", "H"), ("Current Promo (today)", "O"),
        ("Expected Shelf Price Today", "P"), ("Next Promo Start", "Q"), ("Execution Priority / Callouts", "R"),
        ("Data Status", "U"), ("Open Issues", "X")]
ws_sh["A7"] = "PLAN FOR THIS PROGRAM"
ws_sh["A7"].font = f(True, "FFFFFF")
for c in range(1, 12):
    ws_sh.cell(row=7, column=c).fill = fill(WINE)
for i, (label, col) in enumerate(snap):
    r = 8 + i
    ws_sh.cell(row=r, column=1, value=label).font = f(True, SLATE, 9)
    ws_sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    cell = ws_sh.cell(row=r, column=3,
                      value=f'=IF($C$4="","",IFERROR(INDEX(\'Chain Program\'!${col}${CP_FIRST}:${col}${CP_LAST},'
                            f'MATCH($C$4,{CPA},0))&"","Unknown program"))')
    ws_sh.merge_cells(start_row=r, start_column=3, end_row=r, end_column=11)
    cell.alignment = WRAP
    cell.font = f()
    if col == "P":
        cell.value = (f'=IF($C$4="","",IFERROR(INDEX(\'Chain Program\'!$P${CP_FIRST}:$P${CP_LAST},'
                      f'MATCH($C$4,{CPA},0)),"Unknown program"))')
        cell.number_format = '"$"#,##0.00'
        cell.font = f(True, WINE, 12)
    if col == "Q":
        cell.value = (f'=IF($C$4="","",IFERROR(INDEX(\'Chain Program\'!$Q${CP_FIRST}:$Q${CP_LAST},'
                      f'MATCH($C$4,{CPA},0)),"Unknown program"))')
        cell.number_format = "m/d/yyyy"
    if col == "R":
        ws_sh.row_dimensions[r].height = 42

HR = 8 + len(snap) + 1
ws_sh.cell(row=HR, column=1, value="VISIT HISTORY (newest first)").font = f(True, "FFFFFF")
for c in range(1, 12):
    ws_sh.cell(row=HR, column=c).fill = fill(TEAL)
hist_cols = [("#", None, 5), ("Visit Date", "A", 11), ("Rep", "B", 16), ("Store # / Location", "F", 20),
             ("SKU / Pack", "G", 16), ("Observed ($)", "I", 10), ("Price Check", "J", 11),
             ("Discrepancy Type", "O", 22), ("Visit Notes", "P", 44), ("Follow-up Action", "Q", 30), ("Status", "T", 11)]
for ci, (name, _, w) in enumerate(hist_cols, start=1):
    cell = ws_sh.cell(row=HR + 1, column=ci, value=name)
    cell.font = f(True, "FFFFFF", 9)
    cell.fill = fill(NAVY)
    cell.alignment = CENTER
    ws_sh.column_dimensions[get_column_letter(ci)].width = w
for k in range(1, HIST_ROWS + 1):
    r = HR + 1 + k
    ws_sh.cell(row=r, column=1, value=k).font = f(color=SLATE)
    for ci, (name, col, _) in enumerate(hist_cols[1:], start=2):
        cell = ws_sh.cell(row=r, column=ci,
                          value=f'=IFERROR(INDEX(\'Visit Log\'!${col}${CP_FIRST}:${col}${VL_LAST},'
                                f'MATCH($A{r},\'Visit Log\'!$X${CP_FIRST}:$X${VL_LAST},0))&"","")')
        cell.border = BOX
        cell.alignment = WRAP
        cell.font = f(size=9)
    # dates / prices must stay numeric for formatting
    for ci, col, fmt in ((2, "A", "m/d/yyyy"), (6, "I", '"$"#,##0.00')):
        cell = ws_sh.cell(row=r, column=ci)
        cell.value = (f'=IFERROR(IF(INDEX(\'Visit Log\'!${col}${CP_FIRST}:${col}${VL_LAST},'
                      f'MATCH($A{r},\'Visit Log\'!$X${CP_FIRST}:$X${VL_LAST},0))="","",'
                      f'INDEX(\'Visit Log\'!${col}${CP_FIRST}:${col}${VL_LAST},'
                      f'MATCH($A{r},\'Visit Log\'!$X${CP_FIRST}:$X${VL_LAST},0))),"")')
        cell.number_format = fmt
ws_sh.conditional_formatting.add(
    f"K{HR + 2}:K{HR + 1 + HIST_ROWS}", CellIsRule(operator="equal", formula=['"Open"'], fill=fill("FECACA")))
ws_sh.conditional_formatting.add(
    f"G{HR + 2}:G{HR + 1 + HIST_ROWS}", CellIsRule(operator="equal", formula=['"Over plan"'], fill=fill("FECACA")))

# ---------------------------------------------------------------- Dashboard
ws = ws_dash
ws.sheet_view.showGridLines = False
ws["A1"] = "CHAIN AUDIT DASHBOARD"
ws["A1"].font = f(True, NAVY, 16)
ws["A2"] = "=\"As of \"&TEXT(TODAY(),\"m/d/yyyy\")&\"  |  \"&\"" + AS_OF.replace('"', "'") + "\""
ws["A2"].font = f(italic=True, color=SLATE, size=9)
for col, w in zip("ABCDEFG", (30, 14, 4, 30, 14, 14, 14)):
    ws.column_dimensions[col].width = w

kpis = [
    ("Programs tracked", f'=COUNTA({CPA})', "0"),
    ("Programs needing 2026 confirmation", f"=COUNTIF('Chain Program'!$U${CP_FIRST}:$U${CP_LAST},\"Confirm for 2026\")", "0"),
    ("Promo windows active today", f'=COUNTIF({PCL},"Active")', "0"),
    ("Promo windows starting in next 14 days", f'=COUNTIFS({PCI},">"&TODAY(),{PCI},"<="&TODAY()+14)', "0"),
    ("Visits logged, last 30 days", f'=COUNTIFS({VLA},">="&TODAY()-30)', "0"),
    ("Open / in-progress discrepancies", f'=COUNTIF({VLT},"Open")+COUNTIF({VLT},"In Progress")', "0"),
    ("Overdue follow-ups", f'=COUNTIFS({VLT},"Open",{VLS},"<"&TODAY())+COUNTIFS({VLT},"In Progress",{VLS},"<"&TODAY())', "0"),
    ("Shelf price match rate",
     f'=IFERROR(COUNTIF({VLJ},"Match")/(COUNTIF({VLJ},"Match")+COUNTIF({VLJ},"Over plan")+COUNTIF({VLJ},"Under plan")),"No checks yet")',
     "0.0%"),
]
ws["A4"] = "KEY NUMBERS"
ws["A4"].font = f(True, "FFFFFF")
ws["A4"].fill = fill(NAVY)
ws["B4"].fill = fill(NAVY)
for i, (label, formula, fmt) in enumerate(kpis):
    r = 5 + i
    ws.cell(row=r, column=1, value=label).font = f()
    c = ws.cell(row=r, column=2, value=formula)
    c.font = f(True, NAVY, 12)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    for cc in (1, 2):
        ws.cell(row=r, column=cc).border = BOX

# by brand
ws["D4"] = "BY BRAND"
for c in "DEFG":
    ws[f"{c}4"].fill = fill(WINE)
ws["D4"].font = f(True, "FFFFFF")
for ci, h in enumerate(("Brand", "Programs", "Active promos", "Open issues"), start=4):
    cell = ws.cell(row=5, column=ci, value=h)
    cell.font = f(True, size=9)
    cell.fill = fill(LIGHT)
    cell.border = BOX
for i, b in enumerate(pdat.BRANDS):
    r = 6 + i
    ws.cell(row=r, column=4, value=b).font = f()
    ws.cell(row=r, column=5, value=f"=COUNTIF('Chain Program'!$B${CP_FIRST}:$B${CP_LAST},D{r})")
    ws.cell(row=r, column=6, value=f"=COUNTIFS('Promo Calendar'!$B${CP_FIRST}:$B${PC_LAST},D{r},{PCL},\"Active\")")
    ws.cell(row=r, column=7, value=f'=COUNTIFS({VLD},D{r},{VLT},"Open")+COUNTIFS({VLD},D{r},{VLT},"In Progress")')
    for c in range(4, 8):
        ws.cell(row=r, column=c).border = BOX
        if c > 4:
            ws.cell(row=r, column=c).font = f()

# by rep
RR = 15
ws.cell(row=RR, column=1, value="ACCOUNTABILITY BY REP").font = f(True, "FFFFFF")
for c in range(1, 8):
    ws.cell(row=RR, column=c).fill = fill(SLATE)
rep_heads = ("Rep", "Visits (30 days)", "", "Open issues owned", "Overdue", "Last visit", "Price checks (30 days)")
for ci, h in enumerate(rep_heads, start=1):
    if not h:
        continue
    cell = ws.cell(row=RR + 1, column=ci, value=h)
    cell.font = f(True, size=9)
    cell.fill = fill(LIGHT)
    cell.border = BOX
roster_n = 26
for k in range(roster_n):
    r = RR + 2 + k
    lr = 4 + k
    ws.cell(row=r, column=1, value=f'=IF(Lists!$A${lr}="","",Lists!$A${lr})').font = f()
    a = f"$A{r}"
    ws.cell(row=r, column=2, value=f'=IF({a}="","",COUNTIFS({VLB},{a},{VLA},">="&TODAY()-30))')
    ws.cell(row=r, column=4, value=f'=IF({a}="","",COUNTIFS({VLR},{a},{VLT},"Open")+COUNTIFS({VLR},{a},{VLT},"In Progress"))')
    ws.cell(row=r, column=5, value=(f'=IF({a}="","",COUNTIFS({VLR},{a},{VLT},"Open",{VLS},"<"&TODAY())'
                                    f'+COUNTIFS({VLR},{a},{VLT},"In Progress",{VLS},"<"&TODAY()))'))
    ws.cell(row=r, column=6, value=f'=IF({a}="","",IF(COUNTIFS({VLB},{a})=0,"No visits",_xlfn.MAXIFS({VLA},{VLB},{a})))')
    ws.cell(row=r, column=6).number_format = "m/d/yyyy"
    ws.cell(row=r, column=7, value=f'=IF({a}="","",COUNTIFS({VLB},{a},{VLA},">="&TODAY()-30,{VLJ},"?*"))')
    for c in (1, 2, 4, 5, 6, 7):
        ws.cell(row=r, column=c).border = BOX
        ws.cell(row=r, column=c).font = f()
ws.conditional_formatting.add(f"E{RR + 2}:E{RR + 1 + roster_n}",
                              CellIsRule(operator="greaterThan", formula=["0"], fill=fill("FECACA")))

# by discrepancy type
DR = RR + roster_n + 3
ws.cell(row=DR, column=1, value="OPEN DISCREPANCIES BY TYPE").font = f(True, "FFFFFF")
for c in (1, 2):
    ws.cell(row=DR, column=c).fill = fill(TEAL)
for i, t in enumerate([x for x in pdat.DISCREPANCY_TYPES if x != "None"]):
    r = DR + 1 + i
    ws.cell(row=r, column=1, value=t).font = f()
    ws.cell(row=r, column=2, value=f'=COUNTIFS({VLO},A{r},{VLT},"Open")+COUNTIFS({VLO},A{r},{VLT},"In Progress")').font = f()
    for c in (1, 2):
        ws.cell(row=r, column=c).border = BOX

# ---------------------------------------------------------------- Data Review
ws = ws_dr
title(ws, "DATA REVIEW - QUESTIONS FOUND WHILE MERGING THE TWO DOCUMENTS",
      "Resolve these, then update Chain Program / Promo Calendar and flip Data Status to 'Current (2026)'.", 7)
n_confirm = sum(1 for p in pdat.PROGRAMS if p["status"] == pdat.CONFIRM and "HW" in p["source"])
dr_cols = [("#", 5), ("Account", 26), ("Brand", 13), ("Issue", 60), ("Action Needed", 46), ("Owner", 16), ("Resolved?", 10)]
group_band(ws, 4, [("", 1, 7, NAVY)])
headers(ws, 5, dr_cols, lambda i: NAVY)
for i, (acct, brand, issue, action) in enumerate(pdat.DATA_REVIEW, start=1):
    r = 5 + i
    issue = issue.replace("22 Henry's programs", f"{n_confirm} programs")
    for c, v in enumerate((i, acct, brand, issue, action, None, "N"), start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = f(color=BLUE if c >= 6 else "000000")
        cell.alignment = WRAP
        cell.border = BOX
    ws.cell(row=r, column=6).fill = fill(INPUT_FILL)
    ws.cell(row=r, column=7).fill = fill(INPUT_FILL)
add_list_validation(ws, f"={LIST_RANGES['Y/N']}", f"G6:G{5 + len(pdat.DATA_REVIEW)}")
ws.conditional_formatting.add(f"G6:G{5 + len(pdat.DATA_REVIEW)}",
                              CellIsRule(operator="equal", formula=['"Y"'], fill=fill("D1FAE5")))
ws.freeze_panes = "A6"

# ---------------------------------------------------------------- README
ws = ws_readme
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 100
rows = [
    ("title", "CHAIN AUDIT TOOL - Pricing, Placement & Display Compliance"),
    ("sub", AS_OF),
    ("", ""),
    ("h", "WHAT EACH TAB IS FOR"),
    ("Dashboard", "Team KPIs, accountability by rep, open discrepancies by type. Nothing to type."),
    ("Chain Program", "The hot sheet. One row per brand x account: availability, merchandising callout, pricing and "
                      "the expected shelf price today. Manager-owned."),
    ("Promo Calendar", "Every promo / case-cost window with start and end dates. Status (Active / Upcoming / Expired) "
                       "is automatic. Manager-owned."),
    ("Visit Log", "Reps log each store check: observed price, tag, planogram location, in-stock, display, "
                  "discrepancy type, notes and follow-up. Expected price and Price Check fill in automatically."),
    ("Store History", "Before a visit, pick the program (and store #) to see today's plan and the last 20 visit "
                      "notes for that store, newest first."),
    ("Data Review", "Conflicts and gaps found while merging the Non-Alc Chain doc (9/8/26) with the HW Soda hot sheet "
                    "(8/25/25). Work these down, then mark programs 'Current (2026)'."),
    ("Lists", "Dropdown values: team roster, SKUs, discrepancy types, statuses."),
    ("", ""),
    ("h", "COLOR KEY"),
    ("Blue text", "Plan data the manager maintains (Chain Program, Promo Calendar, Lists)."),
    ("Yellow cells", "Cells reps fill in (Visit Log, Store History selectors)."),
    ("Grey cells", "Formulas. Do not type over them."),
    ("Amber 'Confirm for 2026'", "Program still relies on 2025 hot-sheet data. Treat expected prices as provisional."),
    ("", ""),
    ("h", "FIELD WORKFLOW (REPS)"),
    ("1. Before the visit", "Store History tab: pick the program and store #. Read open issues and last notes."),
    ("2. In the store", "Check price vs. Expected Shelf Price, shelf tag, planogram location, all authorized SKUs, "
                        "and the display callout."),
    ("3. Log it", "Visit Log: add one row per program checked (add extra rows for SKU-specific price checks). "
                  "Pick a Discrepancy Type, write notes, set Follow-up Owner, Due Date and Status."),
    ("4. Next visit", "Update the Status / Resolved Date of prior open rows when fixed. Do not delete history."),
    ("", ""),
    ("h", "MANAGER WORKFLOW"),
    ("Add a program", "New row at the bottom of Chain Program. Program ID must be unique (Brand | Account)."),
    ("Add a promo", "New row in Promo Calendar: pick the Program ID, SKU (or All), dates, and Unit Retail if it "
                    "changes shelf price. Leave Unit Retail blank for case-cost-only windows."),
    ("Weekly", "Dashboard: overdue follow-ups and reps with no visits in 30 days. Data Review: close questions."),
    ("", ""),
    ("h", "SHARING & PERMISSIONS (see docs/RECOMMENDATIONS.md in the repo)"),
    ("Plan (view-only)", "Store this file in SharePoint / OneDrive. Share with the team as 'Can view'. Only the "
                         "manager has edit rights to Chain Program and Promo Calendar."),
    ("Visit notes", "Excel cannot give reps 'add rows but not edit the plan' rights inside one file. Recommended: move "
                    "the Visit Log to a Microsoft List (same columns) with 'create items and edit only their own' "
                    "permissions. Until then, a separate shared 'Visit Log' workbook with version history works."),
    ("", ""),
    ("h", "NOTES"),
    ("Expected price", "Uses the lowest active retail promo on the visit date for that SKU (or 'All'), else Base "
                       "Shelf Price. It recalculates if the plan is edited later - for a frozen audit record, move "
                       "the Visit Log to Microsoft Lists."),
    ("History order", "Store History treats lower rows as newer. Keep the Visit Log in date order (sort A-Z by Visit "
                      "Date if needed)."),
    ("Example row", "Visit Log row 6 and the 'EXAMPLE - delete me' roster entry are samples. Delete both before go-live."),
]
r = 1
for kind, text in rows:
    if kind == "title":
        ws.cell(row=r, column=2, value=text).font = f(True, NAVY, 18)
    elif kind == "sub":
        ws.cell(row=r, column=2, value=text).font = f(italic=True, color=SLATE, size=9)
    elif kind == "h":
        ws.cell(row=r, column=2, value=text).font = f(True, "FFFFFF")
        ws.cell(row=r, column=2).fill = fill(NAVY)
        ws.cell(row=r, column=3).fill = fill(NAVY)
    elif kind:
        ws.cell(row=r, column=2, value=kind).font = f(True)
        c = ws.cell(row=r, column=3, value=text)
        c.font = f()
        c.alignment = WRAP
        if kind == "Yellow cells":
            ws.cell(row=r, column=2).fill = fill(INPUT_FILL)
        if kind == "Grey cells":
            ws.cell(row=r, column=2).fill = fill(FORMULA_FILL)
        if kind == "Blue text":
            ws.cell(row=r, column=2).font = f(True, BLUE)
        if kind.startswith("Amber"):
            ws.cell(row=r, column=2).fill = fill("FDE68A")
    r += 1

# tab colors
for sheet, color in ((ws_readme, NAVY), (ws_dash, NAVY), (ws_cp, WINE), (ws_pc, WINE), (ws_vl, "D4A017"),
                     (ws_sh, "D4A017"), (ws_dr, BROWN), (ws_ls, SLATE)):
    sheet.sheet_properties.tabColor = color

for sheet in (ws_cp, ws_pc, ws_vl, ws_dr):
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "4:5"

out = sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).resolve().parent.parent / "Chain_Audit_Tool.xlsx")
wb.calculation.fullCalcOnLoad = True  # openpyxl stores no cached values; Excel computes on open
wb.save(out)
print(f"wrote {out}: {n_prog} programs, {len(promos)} promo windows, {len(pdat.DATA_REVIEW)} data-review items")
