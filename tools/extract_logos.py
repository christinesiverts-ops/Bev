"""Extract brand logos, product photos and chain logos from the source documents.

    python tools/extract_logos.py "Non_Alc_Chain_Document.docx" "HW_Soda_Chain_Activity.pptx"

Writes normalized PNG/JPEG files into app/static/logos/{brands,products,chains}/.
The mapping below reflects the layout of the 9/8/2026 Non-Alc doc and the 8/25/25 HW hot sheet.
Retailer and brand marks belong to their owners; they are used here for internal field identification only.
"""
import io
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageChops
from pptx import Presentation

OUT = Path(__file__).resolve().parent.parent / "app" / "static" / "logos"

# docx media file -> (folder, slug)
DOCX_MAP = {
    "word/media/image3.png": ("brands", "fever-tree"),
    "word/media/image1.png": ("products", "zoa"),
    "word/media/image4.jpg": ("products", "fever-tree"),
    "word/media/image5.png": ("chains", "safeway"),
    "word/media/image13.png": ("chains", "harmons"),
}
# crops from the Naked Life x WFM flyer (image15.png, 1601x900)
FLYER = "word/media/image15.png"
FLYER_CROPS = {("products", "naked-life"): (88, 250, 460, 470), ("chains", "whole-foods-market"): (1296, 66, 1420, 188)}

# HW hot sheet: retailer logos in row order per slide (left column, below the header)
PPTX_CHAINS = {
    1: ["save-mart", "lucky", "raleys", "albertsons", "bi-mart", "bashas", "bevmo"],
    2: ["ck-market", "kroger", "nugget-markets", "stater-bros", "plaid-pantry", "winco-foods"],
    3: ["metropolitan-market", "northwest-grocers", "town-country-markets", "harvest-foods", "ridleys", "rosauers",
        "yokes"],
}


def trim(img: Image.Image) -> Image.Image:
    img = img.convert("RGBA")
    alpha = img.split()[3]
    if alpha.getextrema()[0] < 250:          # transparent background: trim to visible pixels
        box = alpha.point(lambda a: 255 if a > 16 else 0).getbbox()
        return img.crop(box) if box else img
    rgb = img.convert("RGB")
    bg = Image.new("RGB", rgb.size, rgb.getpixel((0, 0)))
    box = ImageChops.difference(rgb, bg).getbbox()
    return img.crop(box) if box else img


def save(img: Image.Image, folder: str, slug: str, max_size=(480, 200)) -> None:
    img = img.convert("RGBA")
    if folder != "products":
        img = trim(img)
    img.thumbnail(max_size if folder != "products" else (720, 480))
    dest = OUT / folder
    dest.mkdir(parents=True, exist_ok=True)
    if folder == "products":
        bg = Image.new("RGB", img.size, "white")
        bg.paste(img, mask=img.split()[3])
        bg.save(dest / f"{slug}.jpg", "JPEG", quality=84, optimize=True)
    else:
        img.save(dest / f"{slug}.png", optimize=True)
    print(f"{folder}/{slug}")


def main(docx_path: str, pptx_path: str) -> None:
    with zipfile.ZipFile(docx_path) as z:
        for member, (folder, slug) in DOCX_MAP.items():
            save(Image.open(io.BytesIO(z.read(member))), folder, slug)
        flyer = Image.open(io.BytesIO(z.read(FLYER)))
        for (folder, slug), box in FLYER_CROPS.items():
            save(flyer.crop(box), folder, slug)

    prs = Presentation(pptx_path)
    for idx, slide in enumerate(prs.slides, start=1):
        pics = sorted([s for s in slide.shapes if s.shape_type == 13 and s.left < 3_000_000 and s.top > 700_000],
                      key=lambda s: s.top)
        for shape, slug in zip(pics, PPTX_CHAINS[idx]):
            save(Image.open(io.BytesIO(shape.image.blob)), "chains", slug)
        if idx == 1:
            badge = [s for s in slide.shapes if s.shape_type == 13 and s.top < 700_000][0]
            save(Image.open(io.BytesIO(badge.image.blob)), "brands", "henrys")
            bottles = [s for s in slide.shapes if s.shape_type == 13 and s.left > 3_000_000 and s.top > 700_000][0]
            save(Image.open(io.BytesIO(bottles.image.blob)), "products", "henrys")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
