import csv
import io

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .. import pricing
from ..common import flash, opt_float, redirect, render, s, today
from ..db import get_db
from ..models import Issue, Program, Store, User, Visit
from ..security import csrf_protect, current_user, log, require_field, require_manager
from ..store_match import programs_for_store

router = APIRouter()

STORE_FIELDS = [("chain", "Chain / Banner"), ("division", "Division (optional)"), ("store_number", "Store #"),
                ("name", "Store name"), ("address", "Address"), ("city", "City"), ("state", "State"), ("zip", "ZIP"),
                ("market_unit", "Market Unit"), ("lat", "Latitude"), ("lng", "Longitude")]
IMPORT_HEADERS = {"chain": "chain", "chain / banner": "chain", "banner": "chain", "division": "division",
                  "store #": "store_number", "store": "store_number", "store number": "store_number",
                  "name": "name", "store name": "name", "address": "address", "city": "city", "state": "state",
                  "zip": "zip", "market unit": "market_unit", "lat": "lat", "latitude": "lat", "lng": "lng",
                  "lon": "lng", "longitude": "lng"}


def _chains(db):
    return db.scalars(select(Program.chain).distinct().order_by(Program.chain)).all()


@router.get("/stores")
def store_list(request: Request, q: str = "", chain: str = "", user: User = Depends(current_user),
               db: Session = Depends(get_db)):
    stmt = select(Store).where(Store.active.is_(True))
    if chain:
        stmt = stmt.where(Store.chain == chain)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(Store.store_number.ilike(like) | Store.name.ilike(like) | Store.city.ilike(like)
                          | Store.address.ilike(like) | Store.chain.ilike(like))
    stores = db.scalars(stmt.order_by(Store.chain, Store.store_number).limit(300)).all()
    ids = [x.id for x in stores]
    last = dict(db.execute(select(Visit.store_id, func.max(Visit.visit_date)).where(Visit.store_id.in_(ids))
                           .group_by(Visit.store_id)).all()) if ids else {}
    opened = dict(db.execute(select(Issue.store_id, func.count(Issue.id)).where(
        Issue.store_id.in_(ids), Issue.status != "Resolved").group_by(Issue.store_id)).all()) if ids else {}
    return render(request, "stores.html", user=user, stores=stores, last=last, opened=opened, chains=_chains(db),
                  f={"q": q, "chain": chain})


@router.get("/stores/new")
def store_new(request: Request, user: User = Depends(require_field), db: Session = Depends(get_db)):
    return render(request, "store_form.html", user=user, st=None, fields=STORE_FIELDS, chains=_chains(db))


def _apply_store(st: Store, form) -> None:
    for name, label in STORE_FIELDS:
        raw = form.get(name)
        if name in ("lat", "lng"):
            try:
                setattr(st, name, opt_float(raw) if s(raw) else None)
            except ValueError as e:
                raise HTTPException(400, f"{label}: {e}")
        else:
            setattr(st, name, s(raw))
    if not st.chain or not st.store_number:
        raise HTTPException(400, "Chain and Store # are required.")


@router.post("/stores/new", dependencies=[Depends(csrf_protect)])
async def store_create(request: Request, user: User = Depends(require_field), db: Session = Depends(get_db)):
    form = await request.form()
    st = Store(created_by_id=user.id, verified=user.is_manager)
    _apply_store(st, form)
    if db.scalar(select(Store).where(Store.chain == st.chain, Store.store_number == st.store_number)):
        raise HTTPException(400, f"{st.chain} #{st.store_number} already exists.")
    db.add(st)
    db.flush()
    log(db, user, "create", "store", st.id, st.label)
    db.commit()
    flash(request, f"Added {st.label}." + ("" if st.verified else " Your manager will verify it."))
    return redirect(f"/stores/{st.id}")


@router.get("/stores/import")
def store_import_form(request: Request, user: User = Depends(require_manager)):
    return render(request, "store_import.html", user=user)


def _rows_from_upload(filename: str, raw: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        all_rows = [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        all_rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    if not all_rows:
        return []
    head = [IMPORT_HEADERS.get(h.lower().strip()) for h in all_rows[0]]
    out = []
    for r in all_rows[1:]:
        rec = {k: (r[i] if i < len(r) else "") for i, k in enumerate(head) if k}
        if any(rec.values()):
            out.append(rec)
    return out


@router.post("/stores/import", dependencies=[Depends(csrf_protect)])
async def store_import(request: Request, user: User = Depends(require_manager), db: Session = Depends(get_db)):
    form = await request.form()
    up = form.get("file")
    if not isinstance(up, UploadFile) and not hasattr(up, "read"):
        raise HTTPException(400, "Choose a CSV or Excel file.")
    raw = await up.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(400, "File is larger than 10 MB.")
    try:
        rows = _rows_from_upload(up.filename or "", raw)
    except Exception as e:  # malformed workbook
        raise HTTPException(400, f"Could not read that file: {e}")
    added = updated = skipped = 0
    for rec in rows:
        chain, num = s(rec.get("chain")), s(rec.get("store_number"))
        if not chain or not num:
            skipped += 1
            continue
        st = db.scalar(select(Store).where(Store.chain == chain, Store.store_number == num))
        if st is None:
            st = Store(chain=chain, store_number=num, created_by_id=user.id)
            db.add(st)
            added += 1
        else:
            updated += 1
        st.verified = True
        for k in ("division", "name", "address", "city", "state", "zip", "market_unit"):
            if k in rec:
                setattr(st, k, s(rec[k]))
        for k in ("lat", "lng"):
            try:
                if s(rec.get(k)):
                    setattr(st, k, opt_float(rec[k]))
            except ValueError:
                pass
    log(db, user, "import", "store", None, f"added={added} updated={updated} skipped={skipped}")
    db.commit()
    flash(request, f"Import done: {added} added, {updated} updated, {skipped} skipped (missing chain or store #).")
    return redirect("/stores")


@router.get("/stores/{sid}")
def store_detail(sid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    st = db.get(Store, sid)
    if not st:
        raise HTTPException(404)
    t = today()
    matched, other = programs_for_store(db, st)
    cards = [{"p": p, "offer": pricing.current_offer_text(p, t), "expected": pricing.expected_price(p, "All", t),
              "next": pricing.next_promo_start(p, t), "sku_prices": pricing.active_sku_prices(p, t)} for p in matched]
    visits = db.scalars(select(Visit).where(Visit.store_id == st.id)
                        .order_by(Visit.checked_in_at.desc()).limit(20)).all()
    issues = db.scalars(select(Issue).where(Issue.store_id == st.id, Issue.status != "Resolved")
                        .order_by(Issue.created_at.desc())).all()
    return render(request, "store_detail.html", user=user, st=st, cards=cards, other=other, visits=visits,
                  issues=issues, t=t)


@router.get("/stores/{sid}/edit")
def store_edit(sid: int, request: Request, user: User = Depends(require_manager), db: Session = Depends(get_db)):
    st = db.get(Store, sid)
    if not st:
        raise HTTPException(404)
    return render(request, "store_form.html", user=user, st=st, fields=STORE_FIELDS, chains=_chains(db))


@router.post("/stores/{sid}/edit", dependencies=[Depends(csrf_protect)])
async def store_update(sid: int, request: Request, user: User = Depends(require_manager),
                       db: Session = Depends(get_db)):
    st = db.get(Store, sid)
    if not st:
        raise HTTPException(404)
    form = await request.form()
    _apply_store(st, form)
    clash = db.scalar(select(Store).where(Store.chain == st.chain, Store.store_number == st.store_number,
                                          Store.id != st.id))
    if clash:
        raise HTTPException(400, f"{st.chain} #{st.store_number} already exists.")
    st.verified = form.get("verified") == "Y"
    st.active = form.get("active", "Y") == "Y"
    log(db, user, "update", "store", st.id, st.label)
    db.commit()
    flash(request, "Store saved.")
    return redirect(f"/stores/{st.id}")
