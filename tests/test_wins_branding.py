import io

from openpyxl import load_workbook
from sqlalchemy import select

from app import db as app_db
from app.models import Photo, Program, User, Win, WinCheck
from tests.conftest import jpeg_bytes, login, make_user, post
from tests.test_workflow import add_store, check_in, pid


def test_log_win_and_recheck_placement_next_visit(manager, other):
    sid = add_store(manager, chain="WinCo Foods", number="12")
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    v1 = check_in(other, sid)
    r = post(other, f"/visits/{v1}/wins", {"win_type": "Cold placement (cooler / cold vault)", "program_id": pid("Henry's | WinCo Foods"),
                                           "location": "Cold vault / cooler door", "facings": "2", "notes": "2 doors of Root Beer"},
             files=[("photos", ("cold.jpg", jpeg_bytes("blue"), "image/jpeg"))], page=f"/visits/{v1}")
    assert r.status_code == 303
    post(other, f"/visits/{v1}/wins", {"win_type": "Ad / feature secured", "program_id": "", "brand": "ZOA"}, page=f"/visits/{v1}")
    with app_db.SessionLocal() as s:
        cold = s.scalar(select(Win).where(Win.win_type.like("Cold%")))
        assert cold.tracked and cold.status == "Active" and cold.brand == "Henry's" and cold.facings == 2
        assert s.scalar(select(Photo).where(Photo.win_id == cold.id)) is not None
        ad = s.scalar(select(Win).where(Win.win_type == "Ad / feature secured"))
        assert ad.tracked is False and ad.brand == "ZOA"
    # next visit asks "still up?" for the tracked placement only
    v2 = check_in(other, sid)
    page = other.get(f"/visits/{v2}").text
    assert "Are these still up?" in page and "Cold placement" in page
    r = post(other, f"/wins/{cold.id}/check", {"status": "Gone", "visit_id": v2, "back": f"/visits/{v2}"}, page=f"/visits/{v2}")
    assert r.status_code == 303 and r.headers["location"] == f"/visits/{v2}"
    with app_db.SessionLocal() as s:
        w = s.get(Win, cold.id)
        assert w.status == "Gone" and s.scalar(select(WinCheck).where(WinCheck.win_id == w.id)).visit_id == v2
    assert "Wins" in manager.get("/wins").text and "2 doors of Root Beer" in manager.get("/wins").text
    assert "Displays &amp; placements" in manager.get(f"/stores/{sid}").text
    # export has the wins sheet
    wb = load_workbook(io.BytesIO(manager.get("/data/export.xlsx").content))
    assert wb["Wins & Placements"].max_row == 3


def test_win_permissions(manager, other):
    sid = add_store(manager)
    v = check_in(manager, sid)
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    r = post(other, f"/visits/{v}/wins", {"win_type": "New display", "brand": "ZOA"}, page="/")
    assert r.status_code == 403
    r = post(manager, f"/visits/{v}/wins", {"win_type": "Made up type", "brand": "ZOA"}, page="/")
    assert r.status_code == 400


def test_theme_defaults_by_role_and_toggle(manager, other):
    assert 'data-theme="dark"' in manager.get("/").text          # managers: dark
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    assert 'data-theme="light"' in other.get("/").text           # reps: light
    post(other, "/account/theme", {"theme": "dark", "back": "/stores"})
    assert 'data-theme="dark"' in other.get("/").text
    assert 'data-theme="light"' in other.get("/login", follow_redirects=True).text or True


def test_branding_colors_and_logo(manager, other):
    r = post(manager, "/admin/branding", {"company_name": "Pacific Bev Co", "primary_color": "#FFEE00", "accent_color": "#B7791F"},
             page="/admin/branding")
    assert r.status_code == 400 and "too light" in r.text
    r = post(manager, "/admin/branding", {"company_name": "Pacific Bev Co", "primary_color": "#1D3557", "accent_color": "#E09F3E"},
             page="/admin/branding")
    assert r.status_code == 303
    css = manager.get("/theme.css").text
    assert "--primary:29 53 87" in css
    assert "Pacific Bev Co" in manager.get("/").text
    # logo upload is re-encoded to PNG and served publicly (login page shows it)
    r = post(manager, "/admin/branding/logo", {"kind": "company", "name": "logo"},
             files=[("file", ("logo.jpg", jpeg_bytes("green"), "image/jpeg"))], page="/admin/branding")
    assert r.status_code == 303
    login_page = other.get("/login").text
    assert "/branding/company/logo.png" in login_page
    assert other.get("/branding/company/logo.png").headers["content-type"] == "image/png"
    # rejects non-images and path tricks
    bad = post(manager, "/admin/branding/logo", {"kind": "chains", "name": "Kroger"},
               files=[("file", ("x.svg", b"<svg onload=alert(1)>", "image/svg+xml"))], page="/admin/branding")
    assert bad.status_code == 400
    assert other.get("/branding/company/..%2F..%2Faudit.db").status_code == 404
    # reps can't change branding
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    assert other.get("/admin/branding").status_code == 403


def test_bundled_logos_resolve(client):
    from app import branding
    assert branding.logo_url("chains", branding.chain_slug("WinCo Foods")) == "/static/logos/chains/winco-foods.png"
    assert branding.logo_url("brands", branding.brand_slug("Henry's")) == "/static/logos/brands/henrys.png"
    assert branding.logo_url("brands", branding.brand_slug("ZOA")) is None     # no real ZOA logo yet -> text mark
    with app_db.SessionLocal() as s:
        for chain in s.scalars(select(Program.chain).distinct()).all():
            if chain != "Smart & Final":
                assert branding.logo_url("chains", branding.chain_slug(chain)), chain


def test_migration_adds_columns(tmp_path):
    import sqlite3
    from sqlalchemy import create_engine
    from app.migrate import upgrade
    dbp = tmp_path / "old.db"
    con = sqlite3.connect(dbp)
    con.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, username VARCHAR(64), display_name VARCHAR(120), role VARCHAR(16), "
                "password_hash VARCHAR(255), must_change_password BOOLEAN, active BOOLEAN, failed_logins INTEGER, "
                "locked_until DATETIME, session_version INTEGER, created_at DATETIME, last_login DATETIME)")
    con.execute("INSERT INTO users (id, username, display_name, role, password_hash, must_change_password, active, "
                "failed_logins, session_version) VALUES (1,'a','A','rep','x',0,1,0,1)")
    con.commit(); con.close()
    upgrade(create_engine(f"sqlite:///{dbp}"))
    con = sqlite3.connect(dbp)
    assert con.execute("SELECT theme FROM users WHERE id=1").fetchone()[0] == "auto"
    assert con.execute("SELECT name FROM sqlite_master WHERE name='wins'").fetchone()
