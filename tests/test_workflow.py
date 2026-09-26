import io
import re

from openpyxl import load_workbook
from sqlalchemy import select

from app import db as app_db
from app.models import AuditLine, Issue, Photo, Program, Store, Task, User, Visit
from tests.conftest import jpeg_bytes, login, make_user, post


def pid(code):
    with app_db.SessionLocal() as s:
        return s.scalar(select(Program.id).where(Program.code == code))


def add_store(client, chain="Smart & Final", number="401", **extra):
    data = {"chain": chain, "store_number": number, "city": "Fresno", "lat": "36.7378", "lng": "-119.7871"} | extra
    r = post(client, "/stores/new", data, page="/stores/new")
    assert r.status_code == 303, r.text
    return int(r.headers["location"].rsplit("/", 1)[1])


def check_in(client, store_id, lat="36.7379", lng="-119.7872"):
    r = post(client, "/visits/new", {"store_id": store_id, "lat": lat, "lng": lng, "accuracy": "12"},
             page=f"/visits/new?store_id={store_id}")
    assert r.status_code == 303, r.text
    return int(r.headers["location"].rsplit("/", 1)[1])


def test_full_rep_visit_flow(manager, other):
    store_id = add_store(manager)                       # manager-added store is verified
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    vid = check_in(other, store_id)
    page = other.get(f"/visits/{vid}")
    assert "At store" in page.text and "Singles 2 for $3" in page.text

    sf = pid("ZOA | Smart & Final")
    r = post(other, f"/visits/{vid}/lines",
             {"program_id": sf, "sku": "Singles (Frosted Grape/Tropical Punch)", "observed_price": "1.79",
              "tag_ok": "N", "planogram_ok": "Y", "in_stock_ok": "Y", "display": "No",
              "discrepancy_type": "None", "notes": "2/$3 tag missing", "action": "Hang tag"},
             files=[("photos", ("shelf.jpg", jpeg_bytes(), "image/jpeg"))], page=f"/visits/{vid}")
    assert r.status_code == 303
    with app_db.SessionLocal() as s:
        line = s.scalar(select(AuditLine).where(AuditLine.visit_id == vid))
        assert (line.expected_price, line.price_check, line.discrepancy_type) == (1.50, "Over plan", "Price - over plan")
        iss = s.scalar(select(Issue).where(Issue.audit_line_id == line.id))
        assert iss.status == "Open" and iss.owner.username == "rep1" and "vs plan $1.50" in iss.description
        ph = s.scalar(select(Photo).where(Photo.audit_line_id == line.id))
        assert ph is not None
        v = s.get(Visit, vid)
        assert v.distance_m is not None and v.distance_m < 50
    # photo served to signed-in users only, re-encoded as JPEG
    img = other.get(f"/photos/{ph.id}")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"

    # the plan changing later does not rewrite the frozen audit record
    with app_db.SessionLocal() as s:
        for pr in s.get(Program, sf).promos:
            pr.unit_retail = 9.99
        s.commit()
        assert s.get(AuditLine, line.id).expected_price == 1.50

    # notes are recalled on the store page and the next visit
    assert "2/$3 tag missing" in manager.get(f"/stores/{store_id}").text
    vid2 = check_in(other, store_id)
    assert "2/$3 tag missing" in other.get(f"/visits/{vid2}").text

    # resolving requires a note, then works
    r = post(other, f"/issues/{iss.id}/update", {"status": "Resolved"}, page=f"/issues/{iss.id}")
    assert r.status_code == 400
    r = post(other, f"/issues/{iss.id}/update", {"status": "Resolved", "note": "Tag hung by SM"}, page=f"/issues/{iss.id}")
    assert r.status_code == 303
    with app_db.SessionLocal() as s:
        i = s.get(Issue, iss.id)
        assert i.status == "Resolved" and i.resolved_at and i.updates[-1].note.startswith("Tag hung")

    post(other, f"/visits/{vid}/checkout", page=f"/visits/{vid}")
    with app_db.SessionLocal() as s:
        assert s.get(Visit, vid).checked_out_at is not None


def test_far_away_checkin_is_flagged(manager):
    store_id = add_store(manager, number="402")
    vid = check_in(manager, store_id, lat="37.7749", lng="-122.4194")   # San Francisco vs Fresno
    assert "km from store" in manager.get(f"/visits/{vid}").text


def test_rep_cannot_touch_plan_or_admin(manager, other):
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    sf = pid("ZOA | Smart & Final")
    assert other.get("/plan").status_code == 200                 # can view the plan
    assert other.get("/plan/new").status_code == 403
    assert other.get(f"/plan/{sf}/edit").status_code == 403
    r = post(other, f"/plan/{sf}/promos", {"offer": "free", "start": "2026-01-01", "end": "2026-12-31", "unit_retail": "0.01"})
    assert r.status_code == 403
    assert other.get("/admin/users").status_code == 403
    assert other.get("/data/export.xlsx").status_code == 403
    assert post(other, f"/plan/{sf}/archive").status_code == 403


def test_rep_cannot_edit_someone_elses_visit(manager, other):
    store_id = add_store(manager)
    make_user("rep1", "rep")
    make_user("rep2", "rep")
    login(other, "rep1", "Field-Pass-99")
    vid = check_in(other, store_id)
    post(other, f"/visits/{vid}/lines", {"program_id": pid("ZOA | Smart & Final"), "sku": "12pk", "observed_price": "17.99"},
         page=f"/visits/{vid}")
    with app_db.SessionLocal() as s:
        lid = s.scalar(select(AuditLine.id).where(AuditLine.visit_id == vid))
    post(other, "/logout")
    login(other, "rep2", "Field-Pass-99")
    assert other.get(f"/visits/{vid}").status_code == 200     # team can read
    assert post(other, f"/visits/{vid}/lines", {"program_id": pid("ZOA | Smart & Final")}, page="/").status_code == 403
    assert post(other, f"/lines/{lid}/delete").status_code == 403
    assert post(other, f"/visits/{vid}/notes", {"notes": "hijack"}).status_code == 403
    assert other.get(f"/lines/{lid}/edit").status_code == 403


def test_viewer_is_read_only(manager, other):
    make_user("view1", "viewer")
    login(other, "view1", "Field-Pass-99")
    assert other.get("/dashboard").status_code == 200
    assert other.get("/visits/new").status_code == 403
    assert post(other, "/stores/new", {"chain": "X", "store_number": "1"}).status_code == 403


def test_rep_added_store_is_unverified(manager, other):
    make_user("rep1", "rep")
    login(other, "rep1", "Field-Pass-99")
    sid = add_store(other, chain="WinCo Foods", number="77")
    with app_db.SessionLocal() as s:
        assert s.get(Store, sid).verified is False
    assert "Henry" in other.get(f"/stores/{sid}").text   # WinCo programs match by chain


def test_division_narrows_programs(manager):
    sid = add_store(manager, chain="Albertsons/Safeway", number="1234", division="Seattle Div. #27")
    page = manager.get(f"/stores/{sid}").text
    assert "Seattle Div. #27" in page
    assert "Other Albertsons/Safeway programs" in page


def test_admin_creates_user_with_temp_password(manager, other):
    r = post(manager, "/admin/users", {"display_name": "Pat Rep", "username": "prep", "role": "rep"}, page="/admin/users")
    assert r.status_code == 303
    page = manager.get("/admin/users").text
    temp = re.search(r'class="big-code">([^<]+)<', page).group(1)
    assert "big-code" not in manager.get("/admin/users").text   # shown once
    r = login(other, "prep", temp)
    assert r.headers["location"] == "/account/password"
    # last manager can't be demoted
    with app_db.SessionLocal() as s:
        boss = s.scalar(select(User).where(User.username == "boss"))
    r = post(manager, f"/admin/users/{boss.id}", {"action": "save", "role": "rep", "active": "Y"}, page="/admin/users")
    assert r.status_code == 400


def test_tasks(manager, other):
    make_user("rep1", "rep")
    with app_db.SessionLocal() as s:
        rid = s.scalar(select(User.id).where(User.username == "rep1"))
    post(manager, "/tasks", {"title": "Check WinCo Tier 2 displays", "assignee_id": str(rid), "due_date": "2026-10-13",
                             "program_id": pid("Henry's | WinCo Foods")}, page="/tasks")
    login(other, "rep1", "Field-Pass-99")
    assert "Check WinCo Tier 2 displays" in other.get("/").text
    with app_db.SessionLocal() as s:
        tid = s.scalar(select(Task.id))
    post(other, f"/tasks/{tid}/done", {"note": "All 3 stores built"}, page="/tasks")
    with app_db.SessionLocal() as s:
        assert s.get(Task, tid).status == "Done"


def test_all_pages_render(manager):
    sid = add_store(manager)
    vid = check_in(manager, sid)
    post(manager, f"/visits/{vid}/lines", {"program_id": pid("ZOA | Smart & Final"), "sku": "12pk",
                                           "observed_price": "16.99", "discrepancy_type": "None"}, page=f"/visits/{vid}")
    post(manager, f"/visits/{vid}/lines", {"program_id": pid("ZOA | Smart & Final"),
                                           "sku": "Singles (Frosted Grape/Tropical Punch)", "observed_price": "1.79"},
         page=f"/visits/{vid}")
    prog = pid("Henry's | WinCo Foods")
    with app_db.SessionLocal() as s:
        iid = s.scalar(select(Issue.id))
        lid = s.scalar(select(AuditLine.id))
        prid = s.get(Program, prog).promos[0].id
    with app_db.SessionLocal() as s:
        first = s.scalar(select(AuditLine).order_by(AuditLine.id))
        assert first.price_check == "No plan price" and first.discrepancy_type == "None"   # 12pk promo ended 9/22
    for path in ("/", "/plan", f"/plan/{prog}", f"/plan/{prog}/edit", "/plan/new", f"/promos/{prid}/edit", "/calendar",
                 "/calendar?status=Expired", "/data-review", "/stores", f"/stores/{sid}", f"/stores/{sid}/edit",
                 "/stores/import", "/visits/new", "/visits/new?q=fresno", f"/visits/{vid}", f"/lines/{lid}/edit",
                 "/visits", "/issues", "/issues?status=all&overdue=1", f"/issues/{iid}", "/issues/new", "/tasks",
                 "/dashboard", "/dashboard?days=7", "/admin/users", "/admin/log", "/data", "/account/password"):
        r = manager.get(path)
        assert r.status_code == 200, (path, r.status_code, r.text[:300])


def test_manager_edits_plan(manager):
    wid = pid("Henry's | WinCo Foods")
    r = post(manager, f"/plan/{wid}/promos", {"sku": "All", "offer_type": "TPR", "offer": "TPR $5.49", "unit_retail": "5.49",
                                              "start": "2026-09-20", "end": "2026-10-05", "source": "test"}, page=f"/plan/{wid}")
    assert r.status_code == 303
    assert "$5.49" in manager.get("/plan?brand=Henry%27s").text
    bad = post(manager, f"/plan/{wid}/promos", {"offer": "x", "start": "2026-10-05", "end": "2026-09-01"}, page=f"/plan/{wid}")
    assert bad.status_code == 400


def test_export_and_reimport_roundtrip(manager):
    sid = add_store(manager)
    check_in(manager, sid)
    r = manager.get("/data/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Chain Program", "Promo Calendar", "Stores", "Visits", "Visit Checks", "Follow-ups", "Tasks", "Data Review"} <= set(wb.sheetnames)
    ws = wb["Chain Program"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Henry's | WinCo Foods":
            row[12].value = 6.79                       # Base Shelf Price
    buf = io.BytesIO()
    wb.save(buf)
    r = post(manager, "/data/import-plan", {"replace_promos": "Y"},
             files=[("file", ("plan.xlsx", buf.getvalue(), "application/octet-stream"))], page="/data")
    assert r.status_code == 303
    with app_db.SessionLocal() as s:
        assert s.scalar(select(Program).where(Program.code == "Henry's | WinCo Foods")).base_price == 6.79
        assert s.query(Program).count() == 42
        from app.models import Promo
        assert s.query(Promo).count() == 81


def test_import_original_workbook(manager):
    from pathlib import Path
    raw = (Path(__file__).resolve().parent.parent / "Chain_Audit_Tool.xlsx").read_bytes()
    r = post(manager, "/data/import-plan", {"replace_promos": "Y"},
             files=[("file", ("Chain_Audit_Tool.xlsx", raw, "application/octet-stream"))], page="/data")
    assert r.status_code == 303
    with app_db.SessionLocal() as s:
        from app.models import Promo
        assert s.query(Program).count() == 42 and s.query(Promo).count() == 81


def test_store_import_csv(manager):
    csv_body = ("Chain,Store #,City,State,Lat,Lng,Division\n"
                "WinCo Foods,12,Portland,OR,45.5,-122.6,\n"
                "WinCo Foods,13,Salem,OR,,,\n"
                ",99,Nowhere,,,,\n")
    r = post(manager, "/stores/import", files=[("file", ("stores.csv", csv_body.encode(), "text/csv"))], page="/stores/import")
    assert r.status_code == 303
    with app_db.SessionLocal() as s:
        assert s.query(Store).count() == 2
    r = post(manager, "/stores/import", files=[("file", ("stores.csv", csv_body.encode(), "text/csv"))], page="/stores/import")
    with app_db.SessionLocal() as s:
        assert s.query(Store).count() == 2        # upsert, no duplicates


def test_bad_photo_rejected(manager):
    sid = add_store(manager)
    vid = check_in(manager, sid)
    r = post(manager, f"/visits/{vid}/notes", {"notes": "x"},
             files=[("photos", ("evil.jpg", b"<?php echo 1; ?>", "image/jpeg"))], page=f"/visits/{vid}")
    assert r.status_code == 400


def test_backup(manager, tmp_path):
    r = post(manager, "/data/backup", page="/data")
    assert r.status_code == 303
    assert list((tmp_path / "backups").glob("audit-*.db"))
